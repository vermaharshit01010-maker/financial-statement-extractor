"""Excel export for explicitly approved historical figures."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from valuation_app.models import ReviewedFigure
from valuation_app.numeric import convert_display_scale


EXPORT_COLUMNS = [
    "metric",
    "original_label",
    "value",
    "unit",
    "currency",
    "reporting_period",
    "statement_type",
    "statement_scope",
    "source_page",
    "source_text",
    "review_status",
    "warnings",
    "unit_scale",
    "per_share_unit",
    "unit_exception",
    "source_statement",
    "printed_page",
    "source_row",
    "source_bbox",
    "source_identity",
    "source_value",
    "positive_magnitude",
    "exact_value_text",
    "exact_source_value_text",
]


def create_excel_workbook(
    figures: list[ReviewedFigure], export_scale: str = "source"
) -> bytes:
    """Create an in-memory workbook with numeric values at the requested scale."""
    rows = []
    for figure in figures:
        row = figure.to_dict()
        scaled_value = convert_display_scale(
            figure.value,
            figure.unit_scale or figure.unit,
            export_scale,
        )
        if figure.value is not None and scaled_value is None:
            raise ValueError(
                f"Cannot export {figure.metric} in {export_scale}: "
                "the source unit scale is unresolved."
            )
        row["value"] = scaled_value
        row["source_value"] = figure.source_value
        row["positive_magnitude"] = figure.positive_magnitude
        row["exact_value_text"] = (
            str(scaled_value) if scaled_value is not None else None
        )
        row["exact_source_value_text"] = (
            str(figure.source_value) if figure.source_value is not None else None
        )
        if export_scale != "source":
            row["unit"] = export_scale
            row["unit_scale"] = export_scale
        row["source_bbox"] = (
            ", ".join(str(coordinate) for coordinate in figure.source_bbox)
            if figure.source_bbox
            else None
        )
        row["warnings"] = "; ".join(figure.warnings)
        rows.append(row)
    dataframe = pd.DataFrame(rows, columns=EXPORT_COLUMNS)

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        dataframe.to_excel(writer, index=False, sheet_name="Reviewed Figures")
        worksheet = writer.sheets["Reviewed Figures"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            width = min(
                max(len(str(cell.value or "")) for cell in column_cells) + 2,
                60,
            )
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
    workbook = load_workbook(BytesIO(output.getvalue()))
    worksheet = workbook["Reviewed Figures"]
    column_indexes = {
        cell.value: cell.column
        for cell in worksheet[1]
    }
    for row_index, figure in enumerate(figures, start=2):
        scaled_value = convert_display_scale(
            figure.value,
            figure.unit_scale or figure.unit,
            export_scale,
        )
        for field_name, number in (
            ("value", scaled_value),
            ("source_value", figure.source_value),
            ("positive_magnitude", figure.positive_magnitude),
        ):
            cell = worksheet.cell(row=row_index, column=column_indexes[field_name])
            cell.value = number
            cell.number_format = '#,##0.################;(#,##0.################);-'
    final_output = BytesIO()
    workbook.save(final_output)
    return final_output.getvalue()
