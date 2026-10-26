from pathlib import Path
from io import BytesIO
import json
from decimal import Decimal

from openpyxl import load_workbook

from valuation_app.excel_export import create_excel_workbook
from valuation_app.models import Candidate, ReviewedFigure
from valuation_app.numeric import format_financial_number
from valuation_app.storage import load_candidates, load_review_state, save_candidates, save_review_state


def test_candidates_and_review_state_are_stored_separately(tmp_path: Path):
    candidate = Candidate(
        candidate_id="abc",
        metric="revenue",
        original_label="Revenue from operations",
        value="100",
        unit="crore",
        currency="INR",
        reporting_period="2025",
        statement_type="income_statement",
        source_page=4,
        source_text="Revenue from operations 100",
        statement_scope="consolidated",
    )
    approved = ReviewedFigure(
        candidate_id="abc",
        metric="revenue",
        original_label="Revenue from operations",
        value="100",
        unit="crore",
        currency="INR",
        reporting_period="2025",
        statement_type="income_statement",
        source_page=4,
        source_text="Revenue from operations 100",
        statement_scope="consolidated",
    )

    save_candidates("document", [candidate], tmp_path)
    save_review_state("document", [approved], {"rejected"}, tmp_path)

    loaded = load_candidates("document", tmp_path)
    loaded_approved, rejected_ids = load_review_state("document", tmp_path)
    assert loaded[0].candidate_id == "abc"
    assert loaded[0].value == Decimal("100")
    assert isinstance(loaded[0].value, Decimal)
    assert loaded_approved[0].candidate_id == "abc"
    assert rejected_ids == {"rejected"}
    assert (
        tmp_path / "candidates" / "document.json"
    ) != (tmp_path / "reviews" / "document.json")


def test_excel_export_contains_approved_figures_and_source_page():
    figure = ReviewedFigure(
        candidate_id="abc",
        metric="revenue",
        original_label="Revenue from operations",
        value="100.25",
        unit="crore",
        currency="INR",
        reporting_period="2025",
        statement_type="income_statement",
        source_page=4,
        source_text="Revenue from operations 100.25",
        statement_scope="consolidated",
        unit_scale="million",
        source_statement="Consolidated Statement of Operations",
        printed_page=42,
        source_bbox=(1.0, 2.0, 3.0, 4.0),
        source_value="100.25",
    )

    workbook = load_workbook(BytesIO(create_excel_workbook([figure])))
    sheet = workbook["Reviewed Figures"]
    rows = list(sheet.iter_rows(values_only=True))

    assert rows[0][0] == "metric"
    assert rows[1][0] == "revenue"
    assert rows[1][2] == 100.25
    assert sheet.cell(row=2, column=3).data_type == "n"
    assert rows[1][8] == 4
    assert rows[1][10] == "approved"
    assert rows[1][12] == "million"
    assert rows[1][16] == 42
    assert rows[1][20] == 100.25
    assert rows[1][22] == "100.25"


def test_financial_number_formatting_preserves_precision_and_negative_sign():
    assert format_financial_number(Decimal("1503149")) == "1,503,149"
    assert format_financial_number(Decimal("-108381")) == "(108,381)"
    assert format_financial_number(Decimal("1.23000")) == "1.23000"
    assert format_financial_number(Decimal("-0.000")) == "0.000"
    assert format_financial_number(None) == "—"


def test_excel_export_preserves_source_decimal_precision_as_numeric():
    figure = ReviewedFigure(
        candidate_id="decimal",
        metric="revenue",
        original_label="Revenue",
        value=Decimal("1.234567890123456789"),
        unit="thousand",
        currency="USD",
        reporting_period="2025",
        statement_type="income_statement",
        source_page=1,
        source_text="Revenue 1.234567890123456789",
        statement_scope="consolidated",
        unit_scale="thousand",
    )
    workbook = load_workbook(BytesIO(create_excel_workbook([figure])))
    sheet = workbook["Reviewed Figures"]

    assert sheet.cell(row=2, column=3).data_type == "n"
    assert sheet.cell(row=2, column=3).value == 1.234567890123457
    assert sheet.cell(row=2, column=23).value == "1.234567890123456789"


def test_candidate_json_round_trip_preserves_decimal_text_exactly(tmp_path: Path):
    candidate = Candidate(
        candidate_id="precise",
        metric="revenue",
        original_label="Revenue",
        value=Decimal("1.234567890123456789"),
        unit="thousand",
        currency="USD",
        reporting_period="2025",
        statement_type="income_statement",
        source_page=2,
        source_text="Revenue 1.234567890123456789",
        statement_scope="consolidated",
    )
    save_candidates("precise", [candidate], tmp_path)
    loaded = load_candidates("precise", tmp_path)[0]

    assert loaded.value == Decimal("1.234567890123456789")


def test_explicit_excel_scale_conversion_does_not_mutate_source_figure():
    figure = ReviewedFigure(
        candidate_id="scale",
        metric="revenue",
        original_label="Revenue",
        value=Decimal("1234.500"),
        unit="thousand",
        currency="USD",
        reporting_period="2025",
        statement_type="income_statement",
        source_page=1,
        source_text="Revenue 1234.500",
        statement_scope="consolidated",
        unit_scale="thousand",
    )
    workbook = load_workbook(
        BytesIO(create_excel_workbook([figure], export_scale="million"))
    )
    sheet = workbook["Reviewed Figures"]

    assert sheet["C2"].value == 1.2345
    assert sheet["M2"].value == "million"
    assert figure.value == Decimal("1234.500")
    assert figure.unit_scale == "thousand"


def test_loads_legacy_candidate_and_marks_ambiguous_rupees_unit_unresolved(tmp_path: Path):
    old_record = {
        "candidate_id": "legacy",
        "metric": "revenue",
        "original_label": "Revenue from operations",
        "value": "100",
        "unit": "rupees",
        "currency": "USD",
        "reporting_period": "2025",
        "statement_type": "income_statement",
        "source_page": 1,
        "source_text": "Revenue from operations 100",
        "statement_scope": "consolidated",
        "warnings": [],
    }
    path = tmp_path / "candidates" / "legacy.json"
    path.parent.mkdir()
    path.write_text(json.dumps([old_record]), encoding="utf-8")

    loaded = load_candidates("legacy", tmp_path)

    assert loaded[0].unit is None
    assert loaded[0].unit_scale is None
    assert loaded[0].currency == "USD"
    assert any("Legacy unit metadata" in warning for warning in loaded[0].warnings)
